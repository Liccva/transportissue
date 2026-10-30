import csv
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .data_loader import (
    CONSUMER_NAMES_ROW,
    DEMAND_ROW,
    FIRST_CONSUMER_COL,
    FIRST_SUPPLIER_ROW,
    SUPPLIER_NAME_COL,
    SUPPLY_COL,
)
from .models import EPSILON, Matrix, TransportTable, format_number

DEFAULT_TEMPLATE_PATH = Path(__file__).parent / "templates" / "transport_table_template.xlsx"

PLAN_SHEET_TITLE = "План перевозок"
COSTS_SHEET_TITLE = "Тарифы"
PLAN_CORNER_TEXT = "Матрица объёмов\nперевозок"
COSTS_CORNER_TEXT = "Матрица коэффициентов\nтранспортных затрат"
TOTAL_COST_LABEL = "Стоимость перевозок"


class _AxisMap:
    def __init__(self, first: int, template_last: int, output_last: int):
        self.first, self.template_last, self.output_last = first, template_last, output_last
        self.middle = first + 1 if template_last - first >= 2 else None

    def style_index(self, index: int) -> int:
        if index <= self.first:
            return index
        if index == self.output_last:
            return self.template_last
        return self.middle or self.first

    def start_side_index(self, index: int) -> int:
        if index <= self.first:
            return index
        return self.middle or self.template_last

    def end_side_index(self, index: int) -> int:
        if index < self.first:
            return index
        if index == self.output_last:
            return self.template_last
        return self.middle or self.first


class ResultWriter:
    @staticmethod
    def save_results(
        table: TransportTable,
        path: str | Path,
        template_path: str | Path | None = None,
        *,
        template_sheet: str | None = None,
        into_template: bool = False,
    ) -> None:
        path = Path(path)
        suffix = path.suffix.lower()
        if into_template and suffix != ".xlsx":
            raise ValueError("Запись результатов во входной файл возможна только в формате .xlsx")
        try:
            if suffix == ".xlsx":
                _save_excel(table, path, template_path, template_sheet, into_template)
            elif suffix == ".csv":
                _save_csv(table, path)
            elif suffix == ".json":
                _save_json(table, path)
            else:
                raise ValueError(f"Неподдерживаемый формат результата «{suffix}»: используйте .xlsx, .csv или .json")
        except PermissionError:
            raise PermissionError(f"Нет доступа к файлу {path}: возможно, он открыт в Excel") from None


def _save_excel(
    table: TransportTable,
    path: Path,
    template_path: str | Path | None,
    template_sheet: str | None,
    into_template: bool,
) -> None:
    if template_path is None:
        template_path, template_sheet = DEFAULT_TEMPLATE_PATH, None
    workbook = load_workbook(template_path)
    template = workbook[template_sheet] if template_sheet else workbook.worksheets[0]
    for title in (PLAN_SHEET_TITLE, COSTS_SHEET_TITLE):
        if title in workbook.sheetnames and workbook[title] is not template:
            workbook.remove(workbook[title])

    row_map = _AxisMap(FIRST_SUPPLIER_ROW, _last_filled(template, FIRST_SUPPLIER_ROW, SUPPLIER_NAME_COL, True),
                       FIRST_SUPPLIER_ROW + len(table.suppliers) - 1)
    col_map = _AxisMap(FIRST_CONSUMER_COL, _last_filled(template, FIRST_CONSUMER_COL, CONSUMER_NAMES_ROW, False),
                       FIRST_CONSUMER_COL + len(table.consumers) - 1)

    plan_sheet = workbook.create_sheet(PLAN_SHEET_TITLE)
    _write_table(plan_sheet, template, row_map, col_map, table, table.supply_volumes, PLAN_CORNER_TEXT, True)
    _write_total_cost(plan_sheet, template, row_map.output_last + 2, table.transportation_costs)
    costs_sheet = workbook.create_sheet(COSTS_SHEET_TITLE)
    _write_table(costs_sheet, template, row_map, col_map, table, table.transport_coefficients, COSTS_CORNER_TEXT, False)

    if not into_template:
        for sheet in workbook.worksheets[:]:
            if sheet not in (plan_sheet, costs_sheet):
                workbook.remove(sheet)
    workbook.active = plan_sheet
    for sheet in workbook.worksheets:
        sheet.sheet_view.tabSelected = sheet is plan_sheet
    workbook.save(path)


def _last_filled(sheet: Worksheet, start: int, fixed: int, by_row: bool) -> int:
    index = start
    while sheet.cell(*((index + 1, fixed) if by_row else (fixed, index + 1))).value not in (None, ""):
        index += 1
    return index


def _write_table(
    sheet: Worksheet,
    template: Worksheet,
    row_map: _AxisMap,
    col_map: _AxisMap,
    table: TransportTable,
    matrix: Matrix,
    corner_text: str,
    hide_zeros: bool,
) -> None:
    header_rows, header_cols = (CONSUMER_NAMES_ROW, DEMAND_ROW), (SUPPLIER_NAME_COL, SUPPLY_COL)
    values: dict[tuple[int, int], object] = {(r, c): template.cell(r, c).value
                                             for r in header_rows for c in header_cols}
    values[CONSUMER_NAMES_ROW, SUPPLIER_NAME_COL] = corner_text
    for j, consumer in enumerate(table.consumers):
        values[CONSUMER_NAMES_ROW, FIRST_CONSUMER_COL + j] = consumer.name
        values[DEMAND_ROW, FIRST_CONSUMER_COL + j] = _plain(consumer.resources)
    for i, supplier in enumerate(table.suppliers):
        row = FIRST_SUPPLIER_ROW + i
        values[row, SUPPLIER_NAME_COL] = supplier.name
        values[row, SUPPLY_COL] = _plain(supplier.resources)
        for j in range(len(table.consumers)):
            if not (hide_zeros and matrix[i, j] <= EPSILON):
                values[row, FIRST_CONSUMER_COL + j] = _plain(matrix[i, j])

    for row in range(1, row_map.output_last + 1):
        for col in range(1, col_map.output_last + 1):
            cell = sheet.cell(row, col, values.get((row, col)))
            style_row, style_col = row_map.style_index(row), col_map.style_index(col)
            cell._style = copy(template.cell(style_row, style_col)._style)
            cell.border = Border(
                left=copy(template.cell(style_row, col_map.start_side_index(col)).border.left),
                right=copy(template.cell(style_row, col_map.end_side_index(col)).border.right),
                top=copy(template.cell(row_map.start_side_index(row), style_col).border.top),
                bottom=copy(template.cell(row_map.end_side_index(row), style_col).border.bottom),
            )
    for row in range(1, row_map.output_last + 1):
        sheet.row_dimensions[row].height = template.row_dimensions[row_map.style_index(row)].height
    for col in range(1, col_map.output_last + 1):
        source = template.column_dimensions[get_column_letter(col_map.style_index(col))]
        if source.customWidth:
            sheet.column_dimensions[get_column_letter(col)].width = source.width
    sheet.sheet_view.zoomScale = template.sheet_view.zoomScale


def _write_total_cost(sheet: Worksheet, template: Worksheet, row: int, cost: float) -> None:
    label_source = template.cell(DEMAND_ROW, SUPPLIER_NAME_COL)
    font = copy(label_source.font)
    font.b = True
    thin = Side(style="thin")
    for col in (SUPPLIER_NAME_COL, SUPPLY_COL, FIRST_CONSUMER_COL):
        source = label_source if col != FIRST_CONSUMER_COL else template.cell(DEMAND_ROW, FIRST_CONSUMER_COL)
        cell = sheet.cell(row, col)
        cell._style = copy(source._style)
        cell.font = font
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.cell(row, SUPPLIER_NAME_COL, TOTAL_COST_LABEL)
    sheet.merge_cells(start_row=row, start_column=SUPPLIER_NAME_COL, end_row=row, end_column=SUPPLY_COL)
    sheet.cell(row, FIRST_CONSUMER_COL, _plain(cost))


def _save_csv(table: TransportTable, path: Path) -> None:
    def number(value: float) -> str:
        return format_number(value).replace(".", ",")

    rows = [
        [PLAN_SHEET_TITLE, "Поставщики", *(c.name for c in table.consumers)],
        ["Потребители", "запасы\\потребление", *(number(c.resources) for c in table.consumers)],
        *(
            [s.name, number(s.resources),
             *(number(v) if (v := table.supply_volumes[i, j]) > EPSILON else "" for j in range(len(table.consumers)))]
            for i, s in enumerate(table.suppliers)
        ),
        [],
        [TOTAL_COST_LABEL, number(table.transportation_costs)],
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        csv.writer(file, delimiter=";").writerows(rows)


def _save_json(table: TransportTable, path: Path) -> None:
    def participants(items, volume_key: str) -> list[dict]:
        return [{"name": p.name, volume_key: _plain(p.resources), "fictive": p is table.fictive} for p in items]

    def matrix(source: Matrix) -> list[list[int | float]]:
        return [[_plain(source[i, j]) for j in range(source.num_cols)] for i in range(source.num_rows)]

    data = {
        "suppliers": participants(table.suppliers, "supply"),
        "consumers": participants(table.consumers, "demand"),
        "costs": matrix(table.transport_coefficients),
        "plan": matrix(table.supply_volumes),
        "total_cost": _plain(table.transportation_costs),
    }
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _plain(value: float) -> int | float:
    return int(round(value)) if abs(value - round(value)) <= EPSILON else value
