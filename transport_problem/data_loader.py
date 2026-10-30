import math
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .errors import DataValidationError
from .models import Consumer, Matrix, Supplier, TransportTable, format_number

CONSUMER_NAMES_ROW = 1
DEMAND_ROW = 2
FIRST_SUPPLIER_ROW = 3
SUPPLIER_NAME_COL = 1
SUPPLY_COL = 2
FIRST_CONSUMER_COL = 3


class DataLoader:
    @staticmethod
    def load_data(path: str | Path, *, sheet: str | None = None) -> TransportTable:
        path = Path(path)
        if path.suffix.lower() != ".xlsx":
            raise DataValidationError([f"Неподдерживаемый формат «{path.suffix}», ожидается .xlsx"], source=path.name)
        return _table_from_grid(read_excel_grid(path, sheet), path.name)


def read_excel_grid(path: Path, sheet: str | None = None) -> list[list[Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            raise DataValidationError(
                [f"Лист «{sheet}» не найден. Листы книги: {', '.join(workbook.sheetnames)}"], source=path.name
            )
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _parse_number(value: Any) -> float | None:
    if isinstance(value, str):
        value = re.sub(r"\s", "", value).replace(",", ".")
        try:
            value = float(value)
        except ValueError:
            return None
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        return float(value)
    return None


def _is_empty(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _check_number(value: Any, where: str, errors: list[str]) -> float:
    number = _parse_number(value)
    if _is_empty(value):
        errors.append(f"{where}: пустое значение")
    elif number is None:
        errors.append(f"{where}: ожидалось число, получено {value!r}")
    elif number < 0:
        errors.append(f"{where}: значение должно быть неотрицательным, получено {format_number(number)}")
    else:
        return number
    return 0.0


def _check_names(names: list[str], kind: str, errors: list[str]) -> None:
    if not names:
        errors.append(f"Не найдено ни одного {kind}")
    errors.extend(
        f"Повторяющееся наименование {kind}: «{name}» (×{count})"
        for name, count in Counter(names).items() if count > 1
    )


def _table_from_grid(grid: list[list[Any]], source: str) -> TransportTable:
    def cell(row: int, col: int) -> Any:
        if row <= len(grid) and col <= len(grid[row - 1]):
            return grid[row - 1][col - 1]
        return None

    def where(row: int, col: int, description: str) -> str:
        return f"{get_column_letter(col)}{row} ({description})"

    def names(row: int, col: int, by_row: bool) -> list[str]:
        result = []
        while not _is_empty(value := cell(row + len(result), col) if by_row else cell(row, col + len(result))):
            result.append(str(value).strip())
        return result

    errors: list[str] = []
    consumer_names = names(CONSUMER_NAMES_ROW, FIRST_CONSUMER_COL, by_row=False)
    supplier_names = names(FIRST_SUPPLIER_ROW, SUPPLIER_NAME_COL, by_row=True)
    if not consumer_names and not supplier_names:
        raise DataValidationError(
            ["Таблица не найдена: наименования потребителей ожидаются с ячейки C1, поставщиков — с A3"],
            source=source,
        )
    _check_names(supplier_names, "поставщика", errors)
    _check_names(consumer_names, "потребителя", errors)

    next_row = FIRST_SUPPLIER_ROW + len(supplier_names)
    next_col = FIRST_CONSUMER_COL + len(consumer_names)
    if any(not _is_empty(cell(next_row, col)) for col in range(SUPPLY_COL, next_col)):
        errors.append(f"A{next_row}: не указано наименование поставщика, хотя в строке {next_row} есть данные")
    if any(not _is_empty(cell(row, next_col)) for row in range(DEMAND_ROW, next_row)):
        letter = get_column_letter(next_col)
        errors.append(f"{letter}1: не указано наименование потребителя, хотя в столбце {letter} есть данные")

    consumers = [
        (name, _check_number(cell(DEMAND_ROW, col), where(DEMAND_ROW, col, f"потребление {name}"), errors))
        for col, name in enumerate(consumer_names, start=FIRST_CONSUMER_COL)
    ]
    suppliers, costs = [], []
    for row, name in enumerate(supplier_names, start=FIRST_SUPPLIER_ROW):
        suppliers.append((name, _check_number(cell(row, SUPPLY_COL), where(row, SUPPLY_COL, f"запасы {name}"), errors)))
        costs.append([
            _check_number(cell(row, col), where(row, col, f"тариф {name} → {consumer}"), errors)
            for col, consumer in enumerate(consumer_names, start=FIRST_CONSUMER_COL)
        ])
    if errors:
        raise DataValidationError(errors, source=source)
    return TransportTable(
        [Supplier(name, volume) for name, volume in suppliers],
        [Consumer(name, volume) for name, volume in consumers],
        Matrix(costs),
    )
