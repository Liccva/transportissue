import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from transport_problem import DataLoader, DataValidationError, ResultWriter, solve_transport_problem
from transport_problem.batch import collect_inputs, output_paths, solve_file, write_summary

from .test_transport_problem import TEMPLATE_PATH, make_table

DATA_DIR = Path(__file__).parent.parent / "data"


class TempDirTestCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def write_xlsx(self, name: str, rows: list[list], **sheets: list[list]) -> Path:
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        for title, sheet_rows in sheets.items():
            sheet = workbook.create_sheet(title)
            for row in sheet_rows:
                sheet.append(row)
        workbook.save(path := self.tmp / name)
        return path


class LoaderTests(TempDirTestCase):
    def test_sample_files(self):
        table = DataLoader.load_data(DATA_DIR / "пример.xlsx")
        solve_transport_problem(table)
        self.assertAlmostEqual(table.transportation_costs, 2490)
        self.assertIs(table.fictive, table.consumers[-1])
        table = DataLoader.load_data(DATA_DIR / "несколько_листов.xlsx", sheet="МКТЗ")
        self.assertEqual(len(table.consumers), 10)

    def test_excel_collects_all_errors(self):
        path = self.write_xlsx("bad.xlsx", [["", "", "B1", "B1"], ["", "", 10, None], ["A1", -5, "x", 1]])
        with self.assertRaises(DataValidationError) as context:
            DataLoader.load_data(path)
        errors = context.exception.errors
        self.assertEqual(len(errors), 4, errors)
        for fragment in ("«B1» (×2)", "D2 (потребление B1)", "B3 (запасы A1)", "C3 (тариф A1 → B1)"):
            self.assertIn(fragment, "\n".join(errors))

    def test_missing_name_with_data(self):
        path = self.write_xlsx("gap.xlsx", [["", "", "B1"], ["", "", 3, 1], ["A1", 3, 1], [None, 5, 2]])
        with self.assertRaises(DataValidationError) as context:
            DataLoader.load_data(path)
        self.assertEqual(context.exception.errors, [
            "A4: не указано наименование поставщика, хотя в строке 4 есть данные",
            "D1: не указано наименование потребителя, хотя в столбце D есть данные",
        ])

    def test_sheet_and_format_errors(self):
        path = self.write_xlsx("book.xlsx", [["Описание"]], Данные=[["", "", "B1"], ["", "", 4], ["A1", 4, 3]])
        self.assertEqual(DataLoader.load_data(path, sheet="Данные").suppliers[0].name, "A1")
        with self.assertRaisesRegex(DataValidationError, "Таблица не найдена"):
            DataLoader.load_data(path)
        with self.assertRaisesRegex(DataValidationError, "Листы книги"):
            DataLoader.load_data(path, sheet="Нет")
        with self.assertRaisesRegex(DataValidationError, "ожидается .xlsx"):
            DataLoader.load_data(self.tmp / "a.csv")


class WriterTests(TempDirTestCase):
    def solved(self):
        table = make_table([10, 5], [7, 4], [[1, 2], [3, 4]])
        solve_transport_problem(table)
        return table

    def test_json_output(self):
        table = self.solved()
        ResultWriter.save_results(table, path := self.tmp / "out.json")
        data = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(data["total_cost"], table.transportation_costs)
        self.assertEqual(data["consumers"][-1], {"name": "Фиктивный", "demand": 4, "fictive": True})
        self.assertEqual(data["plan"], table.supply_volumes.items)

    def test_csv_output(self):
        table = self.solved()
        ResultWriter.save_results(table, path := self.tmp / "out.csv")
        with path.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file, delimiter=";"))
        self.assertEqual(rows[0][2:], ["B1", "B2", "Фиктивный"])
        self.assertEqual(rows[-1], ["Стоимость перевозок", f"{table.transportation_costs:g}"])

    def test_into_template(self):
        path = self.tmp / "book.xlsx"
        path.write_bytes(TEMPLATE_PATH.read_bytes())
        table = DataLoader.load_data(path)
        solve_transport_problem(table)
        for _ in range(2):
            ResultWriter.save_results(table, path, path, into_template=True)
        self.assertEqual(load_workbook(path).sheetnames, ["МКТЗ", "План перевозок", "Тарифы"])
        self.assertEqual(DataLoader.load_data(path).transport_coefficients[0, 0], 3)
        with self.assertRaises(ValueError):
            ResultWriter.save_results(table, self.tmp / "out.json", into_template=True)


class BatchTests(TempDirTestCase):
    def test_collect_inputs_and_output_names(self):
        for name in ("a.xlsx", "b.json", "notes.md", "~$a.xlsx"):
            (self.tmp / name).write_text("", encoding="utf-8")
        self.assertEqual([p.name for p in collect_inputs([str(self.tmp)])], ["a.xlsx"])
        self.assertEqual([p.name for p in collect_inputs([str(self.tmp / "*.xlsx"), str(self.tmp / "a.xlsx")])],
                         ["a.xlsx"])
        names = [p.name for p in output_paths([Path("x/п.xlsx"), Path("y/п.xlsx"), Path("д.xlsx")], self.tmp, ".csv")]
        self.assertEqual(names, ["п_1_решение.csv", "п_2_решение.csv", "д_решение.csv"])

    def test_batch_and_summary(self):
        inputs = collect_inputs([str(DATA_DIR)])
        results = [solve_file(src, out) for src, out in zip(inputs, output_paths(inputs, self.tmp, ".xlsx"))]
        by_name = {r.source.name: r for r in results}
        self.assertTrue(by_name["пример.xlsx"].ok and by_name["пример.xlsx"].output.exists())
        self.assertIn("B4 (запасы A2)", by_name["ошибки.xlsx"].error)
        for suffix in (".xlsx", ".csv", ".json"):
            write_summary(results, self.tmp / f"сводка{suffix}")
        data = json.loads((self.tmp / "сводка.json").read_text(encoding="utf-8"))
        self.assertEqual({row["Статус"] for row in data}, {"OK", "Ошибка"})

    def test_errors_do_not_touch_input(self):
        self.assertIn("не найден", solve_file(self.tmp / "nope.xlsx", self.tmp / "o.xlsx").error)
        source = self.tmp / "task.xlsx"
        source.write_bytes(before := (DATA_DIR / "ошибки.xlsx").read_bytes())
        self.assertFalse(solve_file(source, self.tmp / "o.xlsx", into_input=True).ok)
        self.assertEqual(source.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
