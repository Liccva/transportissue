import contextlib
import io
import random
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from transport_problem import (
    Consumer,
    DataLoader,
    Matrix,
    ResultWriter,
    Storage,
    Supplier,
    TransportTable,
    solve_transport_problem,
)

TEMPLATE_PATH = Path(__file__).parent.parent / "data" / "Транспортная_задача_шаблон_данных.xlsx"


def make_table(supply, demand, costs) -> TransportTable:
    return TransportTable(
        [Supplier(f"A{i + 1}", a) for i, a in enumerate(supply)],
        [Consumer(f"B{j + 1}", b) for j, b in enumerate(demand)],
        Matrix(costs),
    )


def min_cost_flow(supply, demand, costs) -> float:
    m, n = len(supply), len(demand)
    source, sink = m + n, m + n + 1
    graph = [[] for _ in range(m + n + 2)]

    def add_edge(u, v, capacity, cost):
        graph[u].append([v, capacity, cost, len(graph[v])])
        graph[v].append([u, 0, -cost, len(graph[u]) - 1])

    for i, a in enumerate(supply):
        add_edge(source, i, a, 0)
    for j, b in enumerate(demand):
        add_edge(m + j, sink, b, 0)
    for i in range(m):
        for j in range(n):
            add_edge(i, m + j, sum(supply), costs[i][j])

    total = 0
    while True:
        dist = [float("inf")] * len(graph)
        prev = [None] * len(graph)
        dist[source] = 0
        for _ in range(len(graph)):
            updated = False
            for u, edges in enumerate(graph):
                if dist[u] == float("inf"):
                    continue
                for k, (v, capacity, cost, _) in enumerate(edges):
                    if capacity > 0 and dist[u] + cost < dist[v]:
                        dist[v], prev[v] = dist[u] + cost, (u, k)
                        updated = True
            if not updated:
                break
        if dist[sink] == float("inf"):
            return total
        flow, v = float("inf"), sink
        while v != source:
            u, k = prev[v]
            flow, v = min(flow, graph[u][k][1]), u
        v = sink
        while v != source:
            u, k = prev[v]
            graph[u][k][1] -= flow
            graph[v][graph[u][k][3]][1] += flow
            v = u
        total += flow * dist[sink]


class StorageTests(unittest.TestCase):
    def test_str(self):
        self.assertEqual(str(Supplier("A1", 20)), "Поставщик A1 (объём запасов: 20)")
        self.assertEqual(str(Consumer("B1", 15)), "Потребитель B1 (объём потребления: 15)")
        self.assertEqual(str(Supplier("A2", 12.5)), "Поставщик A2 (объём запасов: 12.5)")

    def test_storage_is_abstract(self):
        with self.assertRaises(TypeError):
            Storage("X", 1)

    def test_validation(self):
        for name, resources in [("", 1), ("   ", 1), (None, 1), ("A", -1), ("A", "5"), ("A", True),
                                ("A", float("nan"))]:
            with self.subTest(name=name, resources=resources), self.assertRaises(ValueError):
                Supplier(name, resources)
        self.assertEqual(Consumer("B", 0).resources, 0.0)


class MatrixTests(unittest.TestCase):
    def test_access_and_sizes(self):
        simple_matrix = Matrix([[1, 2], [3, 4]])
        self.assertEqual(simple_matrix[1, 1], 4)
        self.assertEqual(simple_matrix.num_rows, 2)
        self.assertEqual(simple_matrix.num_cols, 2)

    def test_sizes_follow_changes(self):
        matrix = Matrix([[1, 2], [3, 4]])
        matrix.add_row([5, 6])
        matrix.add_column([0, 0, 0])
        self.assertEqual((matrix.num_rows, matrix.num_cols), (3, 3))
        matrix.items.append([1, 1, 1])
        self.assertEqual(matrix.num_rows, 4)

    def test_validation(self):
        for items in ([[1, -2]], [[1, "a"]], [[1, 2], [3]]):
            with self.subTest(items=items), self.assertRaises(ValueError):
                Matrix(items)
        with self.assertRaises(ValueError):
            Matrix([[1]])[0, 0] = -1


class TransportTableTests(unittest.TestCase):
    def test_type_validation(self):
        with self.assertRaises(TypeError):
            TransportTable([Consumer("B1", 1)], [Consumer("B2", 1)], Matrix([[1]]))
        with self.assertRaises(TypeError):
            TransportTable([Supplier("A1", 1)], [Supplier("A2", 1)], Matrix([[1]]))

    def test_dimension_validation(self):
        with self.assertRaises(ValueError):
            make_table([5, 5], [10], [[1]])
        with self.assertRaises(ValueError):
            make_table([10], [5, 5], [[1, 2, 3]])

    def test_fictive_consumer(self):
        table = make_table([30, 20], [10, 15], [[1, 2], [3, 4]])
        self.assertEqual(len(table.consumers), 3)
        self.assertEqual(table.consumers[-1].resources, 25)
        self.assertEqual(table.transport_coefficients.num_cols, 3)
        self.assertEqual([table.transport_coefficients[i, 2] for i in range(2)], [0, 0])
        self.assertEqual((table.supply_volumes.num_rows, table.supply_volumes.num_cols), (2, 3))

    def test_fictive_supplier(self):
        table = make_table([10], [15, 5], [[1, 2]])
        self.assertEqual(len(table.suppliers), 2)
        self.assertEqual(table.suppliers[-1].resources, 10)
        self.assertEqual(table.transport_coefficients.num_rows, 2)
        self.assertEqual((table.supply_volumes.num_rows, table.supply_volumes.num_cols), (2, 2))

    def test_balanced_table_unchanged(self):
        table = make_table([10, 5], [7, 8], [[1, 2], [3, 4]])
        self.assertEqual((len(table.suppliers), len(table.consumers)), (2, 2))

    def test_represent(self):
        table = make_table([10, 5], [7, 8], [[1, 2], [3, 4]])
        solve_transport_problem(table)
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            table.represent()
        text = output.getvalue()
        for fragment in ("A1", "A2", "B1", "B2", "Запасы", "Потребление",
                         f"Суммарная стоимость перевозок: {table.transportation_costs:g}"):
            self.assertIn(fragment, text)


class SolverTests(unittest.TestCase):
    def assert_feasible(self, table: TransportTable):
        volumes = table.supply_volumes
        for i, supplier in enumerate(table.suppliers):
            self.assertAlmostEqual(sum(volumes[i, j] for j in range(len(table.consumers))),
                                   supplier.resources)
        for j, consumer in enumerate(table.consumers):
            self.assertAlmostEqual(sum(volumes[i, j] for i in range(len(table.suppliers))),
                                   consumer.resources)

    def test_textbook_example(self):
        table = make_table(
            [7, 9, 18], [5, 8, 7, 14],
            [[19, 30, 50, 10], [70, 30, 40, 60], [40, 8, 70, 20]],
        )
        self.assertIsNone(solve_transport_problem(table))
        self.assert_feasible(table)
        self.assertAlmostEqual(table.transportation_costs, 743)

    def test_degenerate_problem(self):
        table = make_table([10, 10, 10], [10, 10, 10], [[1, 2, 3], [2, 1, 3], [3, 2, 1]])
        solve_transport_problem(table)
        self.assert_feasible(table)
        self.assertAlmostEqual(table.transportation_costs, 30)

    def test_fractional_volumes(self):
        table = make_table([2.5, 3.5], [1.5, 4.5], [[1, 3], [2, 1]])
        solve_transport_problem(table)
        self.assert_feasible(table)
        self.assertAlmostEqual(table.transportation_costs, 1 * 1.5 + 3 * 1 + 1 * 3.5)

    def test_template_file(self):
        table = DataLoader.load_data(TEMPLATE_PATH)
        self.assertEqual(len(table.suppliers), 11)
        self.assertEqual(len(table.consumers), 10)
        self.assertEqual(str(table.suppliers[0]), "Поставщик A1 (объём запасов: 200)")
        self.assertEqual(table.transport_coefficients[9, 9], 5)
        solve_transport_problem(table)
        self.assert_feasible(table)
        supply = [int(s.resources) for s in table.suppliers]
        demand = [int(c.resources) for c in table.consumers]
        costs = table.transport_coefficients.items
        self.assertAlmostEqual(table.transportation_costs, min_cost_flow(supply, demand, costs))

    def test_random_problems_against_min_cost_flow(self):
        rng = random.Random(2026)
        for case in range(150):
            m, n = rng.randint(1, 7), rng.randint(1, 7)
            supply = [rng.randint(0, 60) for _ in range(m)]
            demand = [rng.randint(1, 60) for _ in range(n)]
            costs = [[rng.randint(0, 9) for _ in range(n)] for _ in range(m)]
            table = make_table(supply, demand, [row[:] for row in costs])
            solve_transport_problem(table)
            with self.subTest(case=case, supply=supply, demand=demand, costs=costs):
                self.assert_feasible(table)
                balanced_supply = [int(s.resources) for s in table.suppliers]
                balanced_demand = [int(c.resources) for c in table.consumers]
                expected = min_cost_flow(balanced_supply, balanced_demand,
                                         table.transport_coefficients.items)
                self.assertAlmostEqual(table.transportation_costs, expected)


class DataLoaderTests(unittest.TestCase):
    def test_custom_file(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Матрица", "Поставщики", "X", "Y"])
        sheet.append(["Потребители", "запасы\\потребление", 10, 20])
        sheet.append(["P", 15, 1, 2])
        sheet.append(["Q", 15, 3, 1])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "data.xlsx"
            workbook.save(path)
            table = DataLoader.load_data(path)
        self.assertEqual([s.name for s in table.suppliers], ["P", "Q"])
        self.assertEqual([c.name for c in table.consumers], ["X", "Y"])
        self.assertEqual(table.transport_coefficients.items, [[1, 2], [3, 1]])
        solve_transport_problem(table)
        self.assertAlmostEqual(table.transportation_costs, 10 * 1 + 5 * 2 + 15 * 1)

    def test_invalid_value(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["", "", "B1"])
        sheet.append(["", "", 10])
        sheet.append(["A1", 10, "abc"])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "data.xlsx"
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "C3"):
                DataLoader.load_data(path)


class ResultWriterTests(unittest.TestCase):
    @staticmethod
    def borders(cell):
        b = cell.border
        return b.left.style, b.right.style, b.top.style, b.bottom.style

    def save_and_load(self, table, template_path=TEMPLATE_PATH):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.xlsx"
            ResultWriter.save_results(table, path, template_path)
            return load_workbook(path)

    def test_template_solution_formatting_matches_input(self):
        table = DataLoader.load_data(TEMPLATE_PATH)
        solve_transport_problem(table)
        workbook = self.save_and_load(table)
        self.assertEqual(workbook.sheetnames, ["План перевозок", "Тарифы"])
        template = load_workbook(TEMPLATE_PATH).worksheets[0]
        costs = workbook["Тарифы"]
        for row, template_row in [*((r, r) for r in range(1, 12)), (12, 4), (13, 12)]:
            for col in range(1, 13):
                with self.subTest(row=row, col=col):
                    out, src = costs.cell(row, col), template.cell(template_row, col)
                    self.assertEqual(self.borders(out), self.borders(src))
                    for attribute in ("fill", "font", "alignment", "number_format"):
                        self.assertEqual(repr(getattr(out, attribute)), repr(getattr(src, attribute)))
        self.assertEqual(costs.column_dimensions["A"].width, template.column_dimensions["A"].width)
        self.assertEqual(costs.row_dimensions[1].height, template.row_dimensions[1].height)

        plan = workbook["План перевозок"]
        self.assertEqual(plan["A13"].value, "Фиктивный")
        self.assertEqual(plan["B13"].value, 1980)
        self.assertEqual(plan["A15"].value, "Стоимость перевозок")
        self.assertEqual(plan["C15"].value, 6410)
        for i in range(11):
            for j in range(10):
                expected = table.supply_volumes[i, j]
                self.assertEqual(plan.cell(3 + i, 3 + j).value, expected or None)

    def test_small_table_keeps_outline(self):
        table = make_table([10, 5], [7, 8, 1], [[1, 2, 3], [3, 4, 5]])
        solve_transport_problem(table)
        sheet = self.save_and_load(table)["План перевозок"]
        self.assertEqual(self.borders(sheet["C3"]), ("thin", None, "thin", None))
        self.assertEqual(self.borders(sheet["D4"]), (None, None, None, None))
        self.assertEqual(self.borders(sheet["E5"]), (None, "thin", None, "thin"))
        self.assertEqual(self.borders(sheet["E1"]), (None, "thin", "thin", "thin"))
        self.assertIsNone(sheet["F1"].value)

    def test_single_cell_table(self):
        table = make_table([4], [4], [[2]])
        solve_transport_problem(table)
        sheet = self.save_and_load(table)["План перевозок"]
        self.assertEqual(sheet["C3"].value, 4)
        self.assertEqual(self.borders(sheet["C3"]), ("thin", "thin", "thin", "thin"))
        self.assertEqual(sheet["C5"].value, 8)


if __name__ == "__main__":
    unittest.main()
